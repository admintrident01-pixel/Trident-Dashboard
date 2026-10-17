import copy, io, unittest, zipfile
import stock, stock_import, stock_reports, test_accounts
from core import Problem

def workbook():
    stream=io.BytesIO()
    with zipfile.ZipFile(stream,'w') as z:
        z.writestr('xl/workbook.xml','<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="Stock" sheetId="1" r:id="rId1"/></sheets></workbook>')
        z.writestr('xl/_rels/workbook.xml.rels','<Relationships><Relationship Id="rId1" Target="worksheets/sheet1.xml"/></Relationships>')
        z.writestr('xl/worksheets/sheet1.xml','<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="1"><c r="C1" t="inlineStr"><is><t>Current SOH</t></is></c></row><row r="2"><c r="A2" t="inlineStr"><is><t>BED-MoH</t></is></c><c r="B2" t="inlineStr"><is><t>Bed</t></is></c><c r="C2" t="inlineStr"><is><t>Old asset 123 / New asset 456</t></is></c><c r="E2"><v>5</v></c></row><row r="3"><c r="A3" t="inlineStr"><is><t>bed-moh</t></is></c><c r="B3" t="inlineStr"><is><t>Bed duplicate source</t></is></c></row><row r="4"><c r="A4" t="inlineStr"><is><t>INCOMPLETE</t></is></c></row></sheetData></worksheet>')
    return stream.getvalue()

class StockTests(unittest.TestCase):
    def setUp(self):
        self.state={'staffSettings':{'technicians':[{'name':'Scott Yearbury','staffType':'technician'},{'name':'Rick Leef','staffType':'technician'},{'name':'Office','staffType':'office'}]}}
        self.state=stock.action(self.state,{'action':'product','code':'BED','description':'Bed'},'admin@example.nz')
        self.pid=next(iter(self.state['stock']['products']))
    def run_action(self,action,**values):
        self.state=stock.action(self.state,{'action':action,'date':'2026-10-14',**values},'admin@example.nz')
    def receive(self,pool='New stock',qty=2):
        self.run_action('receive',productId=self.pid,quantity=qty,pool=pool,client='Pickup source' if pool=='Trident refurb stock' else '',technician='Scott Yearbury' if pool=='Trident refurb stock' else '',reference='PICKUP-PO' if pool=='Trident refurb stock' else '',jobId='PICKUP-JOB' if pool=='Trident refurb stock' else '',oldNumbers=['OLD-'+str(i) for i in range(qty)] if type(qty) is int else [])
        return list(self.state['stock']['receipts'])[-1]
    def register(self,rid,numbers):
        self.run_action('asset',receiptId=rid,receiptRevision=self.state['stock']['receipts'][rid]['revision'],numbers=numbers)
        return [a['id'] for a in self.state['stock']['assets'].values() if a['number'] in numbers]
    def act_assets(self,kind,ids,**values):self.run_action(kind,assetIds=ids,revisions={key:self.state['stock']['assets'][key]['revision'] for key in ids},**values)
    def test_accessable_soh_reconciles_orders_delivery_returns_and_receipts(self):
        self.run_action('receive_stock',productId=self.pid,quantity=4,pool='New stock',numbers=['FREE','ORDERED','DELIVERED','RETURNED'],reference='SUP-PO',client='Supplier')
        ids={a['number']:a['id'] for a in self.state['stock']['assets'].values()}
        self.act_assets('allocate',[ids['ORDERED'],ids['DELIVERED'],ids['RETURNED']],client='Client',reference='CLIENT-PO',jobId='JOB-1')
        self.act_assets('dispatch',[ids['DELIVERED'],ids['RETURNED']])
        self.act_assets('return',[ids['RETURNED']],condition='Ready for stock')
        self.act_assets('relabel',[ids['RETURNED']],number='RETURNED-NEW',notes='New label')
        before=copy.deepcopy(self.state)
        report=stock_reports.tables(self.state,'accessable_soh')
        self.assertEqual({r[2] for r in report[3][2]},{'FREE','RETURNED-NEW'})
        self.assertEqual({r[2]:r[-1] for r in report[4][2]},{'ORDERED':'Not delivered','DELIVERED':'Delivered'})
        self.assertEqual(len(report[5][2]),3)
        self.assertIn('RETURNED',[r[4] for r in report[5][2]])
        self.assertIn('Returned',[r[-1] for r in report[5][2]])
        self.assertEqual(sum(r[5] for r in report[6][2]),4)
        self.assertEqual({r[4] for r in report[6][2]},{'FREE','ORDERED','DELIVERED','RETURNED'})
        self.assertEqual(self.state,before)
        import base64
        from openpyxl import load_workbook
        from pypdf import PdfReader
        excel=stock_reports.create(self.state,{'scope':'accessable_soh','format':'xlsx'})
        book=load_workbook(io.BytesIO(base64.b64decode(excel['file'])))
        self.assertEqual(len(book.worksheets),8)
        self.assertEqual(book.worksheets[4].max_row,3)
        pdf=stock_reports.create(self.state,{'scope':'accessable_soh','format':'pdf'})
        text=' '.join(' '.join(page.extract_text().split()) for page in PdfReader(io.BytesIO(base64.b64decode(pdf['file']))).pages)
        for expected in ['New Accessable SOH report','CLIENT-PO','SUP-PO','Not delivered','Delivered']:self.assertIn(expected,text)
        self.assertEqual(self.state,before)

    def test_accessable_soh_includes_unnumbered_opening_and_unallocated_delivery(self):
        imported=stock_import.apply(self.state,stock_import.preview(workbook()),'admin@example.nz',workbook())
        report=stock_reports.tables(imported,'accessable_soh')
        self.assertEqual(sum(r[3] for r in report[3][2]),5)
        self.assertEqual(report[3][2][0][2],'Not assigned')
        self.assertEqual(report[6][2],[])
        self.run_action('receive_stock',productId=self.pid,quantity=1,pool='New stock',numbers=['NO-ORDER'])
        asset=next(iter(self.state['stock']['assets'].values()))
        asset.update(status='With client',reference='',allocationId='')
        report=stock_reports.tables(self.state,'accessable_soh')
        self.assertEqual(report[3][2][0][-1],'Delivered')
        self.assertEqual(report[4][2],[])

    def test_shareable_report_all_assets_low_stock_and_missing_quantities(self):
        self.run_action('receive_stock',productId=self.pid,quantity=5,pool='New stock',numbers=['FREE','ALLOC','DELIVERED','REPAIR','RETIRED'])
        ids={a['number']:a['id'] for a in self.state['stock']['assets'].values()}
        self.act_assets('allocate',[ids['ALLOC'],ids['DELIVERED']],client='Client',reference='PO-1',jobId='JOB-1')
        self.act_assets('dispatch',[ids['DELIVERED']])
        self.act_assets('repair',[ids['REPAIR']],notes='Service')
        self.act_assets('retire',[ids['RETIRED']],notes='Retired')
        self.run_action('stock_settings',minimums={self.pid:2})
        data=stock_reports.tables(self.state,'accessable_soh')
        summary=data[1][2][0]
        self.assertEqual(summary[2:5],[1,1,1]);self.assertEqual(summary[6],'Low stock')
        self.assertEqual({r[2] for r in data[2][2]},set(ids))
        self.assertEqual(sum(r[3] for r in data[2][2]),5)
        self.assertIn('Retired - not in stock',[r[4] for r in data[2][2]])
        self.act_assets('allocate',[ids['FREE']],client='Client',reference='PO-2',jobId='JOB-2')
        self.assertEqual(stock_reports.tables(self.state,'accessable_soh')[1][2][0][6],'Out of stock')
        imported=stock_import.apply(self.state,stock_import.preview(workbook()),'admin',workbook())
        imported=stock.action(imported,{'action':'stock_settings','minimums':{key:6 for key in imported['stock']['products']}},'admin')
        data=stock_reports.tables(imported,'accessable_soh')
        missing=next(r for r in data[1][2] if r[1]=='Bed duplicate source')
        self.assertEqual(missing[2],'Not recorded');self.assertEqual(missing[6],'Quantity needs review')
        confirmed=next(r for r in data[1][2] if r[0]=='BED-MoH')
        self.assertEqual(confirmed[6],'Low stock')
        for invalid in [-1,1.5,True,'2',10001]:
            with self.assertRaises(Problem):self.run_action('stock_settings',minimums={self.pid:invalid})

    def test_minimum_settings_are_per_condition_atomic_and_drive_saved_reports(self):
        self.run_action('receive_stock',productId=self.pid,quantity=2,pool='New stock',numbers=['MIN-A','MIN-B'])
        self.run_action('product',code=self.state['stock']['products'][self.pid]['code'],description='Refurb variant',condition='REFURB')
        other=next(key for key in self.state['stock']['products'] if key!=self.pid)
        self.run_action('stock_settings',minimums={self.pid:3,other:0})
        data=stock.read(self.state)
        self.assertEqual(data['products'][self.pid]['minimumStock'],3)
        self.assertEqual(data['products'][other]['minimumStock'],0)
        row=next(row for row in stock_reports.tables(self.state,'accessable_soh')[1][2] if row[1]!='Refurb variant [REFURB]')
        self.assertEqual(row[6],'Low stock');self.assertEqual(row[7],'Minimum: 3\nShortfall: 1')
        before=copy.deepcopy(self.state)
        for levels in [{self.pid:5,other:-1},{self.pid:5,'missing':1},{self.pid:True},[]]:
            with self.assertRaises(Problem):self.run_action('stock_settings',minimums=levels)
            self.assertEqual(self.state,before)
        self.run_action('stock_settings',minimums={self.pid:2})
        self.assertEqual(stock_reports.tables(self.state,'accessable_soh')[1][2][-1][6],'Stock available')
        self.run_action('stock_settings',minimums={self.pid:None})
        row=next(row for row in stock_reports.tables(self.state,'accessable_soh')[1][2] if row[1]!='Refurb variant [REFURB]')
        self.assertEqual(row[6],'Minimum not set');self.assertEqual(row[7],'Not set')
        self.assertEqual(len(self.state['stock']['assets']),2)

    def test_source_reference_sheet_has_valid_excel_name(self):
        from openpyxl import load_workbook
        raw=stock_reports.excel([('Recorded source PO / SO references',['Reference'],[['PO-1']])])
        sheet=load_workbook(io.BytesIO(raw)).active
        self.assertEqual(sheet['A2'].value,'PO-1')
        self.assertNotIn('/',sheet.title)

    def test_combined_refurb_report_tracks_work_and_number_assignment(self):
        self.run_action('receive_stock',productId=self.pid,quantity=1,pool='New stock',numbers=['BOOKED'])
        ids=list(self.state['stock']['assets'])
        self.act_assets('custom_move',ids,destination='refurbing',technician='Rick Leef',notes='Book workshop')
        self.assertEqual(len(stock_reports.tables(self.state,'refurbing')[0][2]),1)
        self.act_assets('custom_move',ids,destination='refurbed',technician='Rick Leef',notes='Work completed, waiting for label')
        self.assertEqual(stock.read(self.state)['summary']['Available stock'],0)
        report=stock_reports.tables(self.state,'refurbing')
        self.assertEqual(report[0][0],'Refurbishment');self.assertEqual(report[0][2][0][2],'Awaiting new asset number')
        self.act_assets('relabel',ids,number='READY',notes='New label fitted')
        self.assertEqual(stock_reports.tables(self.state,'refurbing')[0][2],[])
        self.assertEqual(stock.read(self.state)['summary']['Available stock'],1)
        self.assertEqual(stock.read(self.state)['assets'][ids[0]]['previousNumbers'],['BOOKED'])

    def test_register_past_numbers_dates_and_reference_without_extra_stock(self):
        rid=self.receive(qty=2)
        rows=[{'currentNumber':'CURRENT-A','number':'OLD-A','validFrom':'2020-01-01','validTo':'2024-01-01','addedAt':'2019-12-01','referenceType':'PO','reference':'HISTORY-PO'},{'currentNumber':'CURRENT-A','number':'OLD-A2','validFrom':'2024-01-01','validTo':'2026-10-14','addedAt':'2024-01-01','referenceType':'SO','reference':'HISTORY-SO'}]
        self.run_action('asset',receiptId=rid,receiptRevision=1,numbers=['CURRENT-A','CURRENT-B'],pastNumbers=rows,assetAddedAt='2019-12-01',labelReferenceType='Reference',labelReference='TRANSFER-REF')
        data=stock.read(self.state);a=next(a for a in data['assets'].values() if a['number']=='CURRENT-A')
        self.assertEqual(data['summary']['Available'],2);self.assertEqual(a['previousNumbers'],['OLD-A','OLD-A2'])
        self.assertEqual([h['number'] for h in a['numberHistory']],['OLD-A','OLD-A2','CURRENT-A'])
        self.assertEqual(a['numberHistory'][0]['reference'],'HISTORY-PO')
        self.assertEqual(a['numberHistory'][-1]['addedAt'],'2019-12-01')
        self.assertEqual(a['numberHistory'][-1]['reference'],'TRANSFER-REF')
        report=stock_reports.tables(self.state,'accessable_soh')[-1]
        self.assertIn('HISTORY-PO',str(report));self.assertIn('2019-12-01',str(report))
        self.assertIn('HISTORY-SO',str(stock_reports.tables(self.state,'asset_history')))

    def test_add_past_history_keeps_current_number_orders_and_stock(self):
        rid=self.receive(qty=1);ids=self.register(rid,['CURRENT'])
        self.act_assets('allocate',ids,client='Client',reference='PO-1',jobId='JOB-1')
        original=copy.deepcopy(self.state['stock']['assets'][ids[0]])
        self.act_assets('past_history',ids,pastNumbers=[{'currentNumber':'CURRENT','number':'OLD','validFrom':'2020-01-01','validTo':'2026-10-14','addedAt':'2020-01-01','referenceType':'Reference','reference':'LEGACY-REF'}])
        a=self.state['stock']['assets'][ids[0]]
        for key in ['number','status','client','reference','jobId','allocationId']:self.assertEqual(a[key],original[key])
        self.assertEqual(a['revision'],original['revision']+1)
        self.assertEqual(stock.read(self.state)['summary']['Allocated'],1)
        self.assertEqual(a['previousNumbers'],['OLD']);self.assertEqual(a['numberHistory'][-1]['number'],'CURRENT')
        before=copy.deepcopy(self.state)
        with self.assertRaises(Problem):self.run_action('past_history',assetIds=ids,revisions={ids[0]:original['revision']},pastNumbers=[])
        self.assertEqual(self.state,before)

    def test_past_history_rejects_invalid_duplicate_overlap_and_wrong_item_atomically(self):
        rid=self.receive(qty=2)
        valid={'currentNumber':'NEW','number':'OLD','validFrom':'2020-01-01','validTo':'2026-10-14','referenceType':'PO','reference':'PO-OLD'}
        cases=[[{**valid,'validTo':'2026-10-15'}],[{**valid,'validFrom':'2026-10-14'}],[{**valid,'validFrom':'bad'}],[{**valid,'addedAt':'bad'}],[{**valid,'referenceType':'Other'}],[{**valid,'currentNumber':'OTHER'}],[{**valid,'number':'new'}],[valid,{**valid,'number':'OTHER-OLD'}],[valid,{**valid,'number':'old'}]]
        for rows in cases:
            before=copy.deepcopy(self.state)
            with self.assertRaises(Problem):self.run_action('asset',receiptId=rid,receiptRevision=1,numbers=['NEW'],pastNumbers=rows)
            self.assertEqual(self.state,before)
        self.run_action('asset',receiptId=rid,receiptRevision=1,numbers=['NEW'],pastNumbers=[valid])
        asset=next(iter(self.state['stock']['assets'].values()));before=copy.deepcopy(self.state)
        with self.assertRaises(Problem):self.act_assets('past_history',[asset['id']],pastNumbers=[{**valid,'number':'OVERLAP'}])
        self.assertEqual(self.state,before)
        with self.assertRaises(Problem):self.run_action('asset',receiptId=rid,receiptRevision=2,numbers=['OLD'])
        self.assertEqual(self.state,before)

    def test_collected_old_label_and_completed_refurb_keep_history_references(self):
        self.run_action('receive',productId=self.pid,quantity=1,pool='Trident refurb stock',client='Client',reference='PICKUP-PO',jobId='PICKUP-JOB',oldNumbers=['COLLECTED-OLD'],oldNumberValidFrom='2020-01-01',oldAssetAddedAt='2019-01-01',oldLabelReferenceType='SO',oldLabelReference='OLD-SO')
        a=next(iter(self.state['stock']['assets'].values()))
        self.assertEqual(a['numberHistory'][0]['reference'],'OLD-SO')
        self.assertEqual(a['numberHistory'][0]['addedAt'],'2019-01-01')
        self.act_assets('custom_move',[a['id']],destination='assets',completeRefurb=True,newNumbers={a['id']:'COMPLETED-NEW'},technician='Rick Leef',notes='Checked and refurbished',assetAddedAt='2026-10-14',labelReferenceType='Reference',labelReference='REFURB-REF')
        current=self.state['stock']['assets'][a['id']]['numberHistory'][-1]
        self.assertEqual(current['reference'],'REFURB-REF');self.assertEqual(current['addedAt'],'2026-10-14')
        self.assertEqual(self.state['stock']['assets'][a['id']]['numberHistory'][0]['reference'],'OLD-SO')

    def test_receive_allocate_dispatch_return_preserves_job_history(self):
        rid=self.receive();ids=self.register(rid,['ASSET-01','ASSET-02'])
        self.assertEqual(stock.read(self.state)['summary']['Available'],2)
        self.act_assets('allocate',ids[:1],client='Client example',reference='PO-100',jobId='TRI-123')
        group=next(iter(self.state['stock']['allocations'].values()));self.assertEqual(group['jobId'],'TRI-123')
        self.act_assets('dispatch',ids[:1]);self.assertEqual(self.state['stock']['assets'][ids[0]]['jobId'],'TRI-123')
        self.act_assets('return',ids[:1],condition='Needs work')
        self.assertEqual(self.state['stock']['assets'][ids[0]]['status'],'Repair')
        self.assertEqual(next(iter(self.state['stock']['allocations'].values()))['assetStates'][ids[0]]['status'],'Returned')
        self.assertEqual(self.state['stock']['movements'][-1]['jobId'],'TRI-123')
        self.act_assets('ready',ids[:1],notes='Checked and repaired')
        self.assertEqual(stock.read(self.state)['summary']['Available'],2)
    def test_cancelled_allocation_does_not_claim_a_client_return(self):
        rid=self.receive(qty=1);ids=self.register(rid,['CANCEL-01'])
        self.act_assets('allocate',ids,client='Client',reference='PO-1',jobId='JOB-1')
        self.act_assets('return',ids,condition='Ready for stock')
        saved=next(iter(self.state['stock']['allocations'].values()))['assetStates'][ids[0]]
        self.assertEqual(saved['status'],'Cancelled');self.assertNotIn('returnedAt',saved)
        self.assertEqual(self.state['stock']['movements'][-1]['kind'],'Allocation cancelled')
        self.assertEqual(stock.read(self.state)['summary']['Available'],1)
    def test_refurb_complete_before_asset_labelling(self):
        rid=self.receive('Trident refurb stock')
        self.run_action('refurb_receipt_ready',receiptId=rid,receiptRevision=1,technician='Rick Leef',notes='Cleaned, replaced cover and checked',date='2026-10-14')
        self.assertEqual(stock.read(self.state)['summary']['Available'],0)
        ids=self.state['stock']['receipts'][rid]['assetIds']
        for i,key in enumerate(ids):self.act_assets('relabel',[key],number='REFURB-'+str(i),notes='New refurb asset label')
        for key in ids:
            asset=self.state['stock']['assets'][key];self.assertEqual(asset['pool'],'Trident refurb stock');self.assertEqual(asset['status'],'Available');self.assertEqual(asset['refurbTechnician'],'Rick Leef');self.assertEqual(asset['refurbCompletedAt'],'2026-10-14')
    def test_refurb_technician_assignment_and_no_early_allocation(self):
        rid=self.receive('Trident refurb stock',1);ids=self.state['stock']['receipts'][rid]['assetIds']
        before=copy.deepcopy(self.state)
        with self.assertRaises(Problem):self.act_assets('allocate',ids,client='Client',reference='PO-1',jobId='JOB-1')
        self.assertEqual(self.state,before)
        self.run_action('assign_refurb',receiptId=rid,receiptRevision=1,technician='Rick Leef')
        self.assertEqual(self.state['stock']['assets'][ids[0]]['refurbTechnician'],'Rick Leef')
        self.act_assets('ready',ids,notes='Ready for use',technician='Rick Leef')
        self.assertEqual(self.state['stock']['assets'][ids[0]]['refurbCompletedAt'],'2026-10-14')
        self.assertEqual(self.state['stock']['assets'][ids[0]]['status'],'Awaiting new asset number')
        self.act_assets('relabel',ids,number='NEW-01',notes='Refurb label replaced')
        self.assertEqual(self.state['stock']['assets'][ids[0]]['status'],'Available')
    def test_pickup_existing_delivered_asset_keeps_identity_and_job_history(self):
        rid=self.receive(qty=1);ids=self.register(rid,['SUPPLIED-01'])
        self.act_assets('allocate',ids,client='Client',reference='DELIVERY-PO',jobId='DELIVERY-JOB')
        self.act_assets('dispatch',ids)
        self.assertEqual(self.state['stock']['assets'][ids[0]]['dispatchedAt'],'2026-10-14')
        self.run_action('receive',productId=self.pid,quantity=1,pool='Trident refurb stock',client='Client',reference='COLLECTION-PO',jobId='COLLECTION-JOB',oldNumbers=['SUPPLIED-01'])
        self.assertEqual(len(self.state['stock']['assets']),1)
        self.assertEqual(self.state['stock']['assets'][ids[0]]['pickupJobId'],'COLLECTION-JOB')
        group=next(iter(self.state['stock']['allocations'].values()))
        self.assertEqual(group['jobId'],'DELIVERY-JOB');self.assertEqual(group['assetStates'][ids[0]]['status'],'Returned')
        self.act_assets('relabel',ids,number='REFURBED-01',notes='New label')
        self.assertEqual(self.state['stock']['assets'][ids[0]]['status'],'Refurbishing')
        self.act_assets('ready',ids,technician='Rick Leef',notes='Work checked')
        self.assertEqual(stock.read(self.state)['summary']['Available'],1)

    def test_pickup_requires_collection_context_and_numbers(self):
        base={'action':'receive','productId':self.pid,'quantity':1,'pool':'Trident refurb stock','client':'Client','reference':'PO','jobId':'JOB','oldNumbers':['OLD']}
        for field in ('client','reference','jobId','oldNumbers'):
            p=base.copy();p.pop(field)
            with self.assertRaises(Problem):stock.action(self.state,p,'admin')

    def test_asset_unique_capacity_versions_and_old_labels(self):
        rid=self.receive(qty=1)
        with self.assertRaises(Problem):self.register(rid,['A','B'])
        ids=self.register(rid,['A'])
        self.act_assets('relabel',ids,number='B',notes='Label replaced')
        self.assertEqual(self.state['stock']['assets'][ids[0]]['previousNumbers'],['A'])
        second=self.receive(qty=1)
        with self.assertRaises(Problem):self.register(second,['A'])
        with self.assertRaises(Problem):self.run_action('allocate',assetIds=ids,revisions={ids[0]:1},client='Client',reference='PO-1',jobId='JOB-1')
        with self.assertRaises(Problem):self.act_assets('allocate',ids,client='Client',reference='PO-1',jobId='')
    def test_complete_refurb_requires_new_number_even_for_legacy_records(self):
        rid=self.receive('Trident refurb stock',1);ids=self.state['stock']['receipts'][rid]['assetIds']
        a=self.state['stock']['assets'][ids[0]];a['requiresRelabel']=False
        before=copy.deepcopy(self.state)
        with self.assertRaises(Problem):self.act_assets('custom_move',ids,destination='assets',completeRefurb=True,technician='Rick Leef',notes='Complete work')
        self.assertEqual(self.state,before)
        self.act_assets('custom_move',ids,destination='assets',completeRefurb=True,technician='Rick Leef',newNumbers={ids[0]:'REFURB-NEW-FINAL'},notes='Complete work and apply label')
        a=stock.read(self.state)['assets'][ids[0]]
        self.assertEqual(a['status'],'Available');self.assertEqual(a['number'],'REFURB-NEW-FINAL')
        self.assertEqual(a['previousNumbers'],['OLD-0'])

    def test_receive_numbered_stock_is_atomic_and_moves_directly_to_on_hand(self):
        self.run_action('receive_stock',productId=self.pid,pool='New stock',quantity=2,numbers=['RECEIVE-1','RECEIVE-2'],client='Supplier',referenceType='PO',reference='SUPPLIER-PO')
        data=stock.read(self.state);row=data['products'][self.pid]
        self.assertEqual((row['incomingQuantity'],row['onHandQuantity']),(0,2))
        self.assertEqual(len(data['receipts']),1);self.assertEqual(len(data['assets']),2)
        self.assertEqual(data['summary']['Awaiting asset numbers'],0)
        for asset in data['assets'].values():self.assertEqual(asset['numberHistory'][0]['validFrom'],'2026-10-14')
        for numbers in [[],['ONE'],['DUPLICATE','duplicate'],['RECEIVE-1','FRESH'],['','FRESH']]:
            before=copy.deepcopy(self.state)
            with self.assertRaises(Problem):self.run_action('receive_stock',productId=self.pid,pool='New stock',quantity=2,numbers=numbers)
            self.assertEqual(self.state,before)

    def test_custom_moves_preserve_identity_numbers_and_client_history(self):
        rid=self.receive(qty=1);ids=self.register(rid,['CUSTOM-OLD'])
        self.act_assets('custom_move',ids,destination='intake',notes='Receiving review')
        self.assertEqual(stock.read(self.state)['products'][self.pid]['incomingQuantity'],1)
        self.act_assets('custom_move',ids,destination='refurbing',technician='Rick Leef',notes='Allocate workshop work')
        self.act_assets('custom_move',ids,destination='refurbed',technician='Rick Leef',notes='Work checked')
        before=copy.deepcopy(self.state)
        with self.assertRaises(Problem):self.act_assets('custom_move',ids,destination='assets',notes='Complete refurb')
        self.assertEqual(self.state,before)
        self.act_assets('custom_move',ids,destination='assets',newNumbers={ids[0]:'CUSTOM-NEW'},notes='Complete refurb and label')
        self.act_assets('custom_move',ids,destination='delivered',client='Client',referenceType='SO',reference='SO-CUSTOM',jobId='JOB-CUSTOM',notes='Direct delivery')
        a=stock.read(self.state)['assets'][ids[0]]
        self.assertEqual(a['status'],'With client');self.assertEqual(a['previousNumbers'],['CUSTOM-OLD'])
        self.assertEqual(a['clientHistory'][0]['referenceType'],'SO')
        self.act_assets('custom_move',ids,destination='intake',notes='Returned for receiving inspection')
        a=stock.read(self.state)['assets'][ids[0]]
        self.assertEqual(a['clientHistory'][0]['status'],'Returned');self.assertEqual(a['status'],'Incoming')
        self.assertEqual(len(self.state['stock']['assets']),1)

    def test_custom_receiving_does_not_bypass_refurb_readiness(self):
        rid=self.receive('Trident refurb stock',1);ids=self.state['stock']['receipts'][rid]['assetIds']
        self.act_assets('custom_move',ids,destination='intake',notes='Hold at receiving')
        before=copy.deepcopy(self.state)
        with self.assertRaises(Problem):self.act_assets('custom_move',ids,destination='assets',newNumbers={ids[0]:'WRONG-READY'},notes='Still needs work')
        self.assertEqual(self.state,before)
        self.act_assets('custom_move',ids,destination='assets',technician='Rick Leef',newNumbers={ids[0]:'VERIFIED-READY'},notes='Refurb work and readiness checked')
        a=stock.read(self.state)['assets'][ids[0]]
        self.assertEqual(a['status'],'Available');self.assertEqual(a['refurbCompletedAt'],'2026-10-14')
        self.assertEqual(self.state['stock']['movements'][-1]['toStatus'],'Available')

    def test_custom_move_batch_is_atomic_and_requires_reason_and_destination_fields(self):
        rid=self.receive(qty=2);ids=self.register(rid,['MOVE-1','MOVE-2'])
        for values in [{'destination':'assets','notes':'No-op'},{'destination':'refurbing','notes':'No technician'},{'destination':'delivered','notes':'No client details'},{'destination':'intake','notes':''}]:
            before=copy.deepcopy(self.state)
            with self.assertRaises(Problem):self.act_assets('custom_move',ids,**values)
            self.assertEqual(self.state,before)
        self.act_assets('allocate',ids[:1],client='Client',reference='PO',jobId='JOB')
        self.act_assets('custom_move',ids[:1],destination='delivered',notes='Deliver existing allocation')
        self.assertEqual(len(self.state['stock']['allocations']),1)
        self.assertEqual(self.state['stock']['assets'][ids[0]]['jobId'],'JOB')

    def test_product_overview_counts_match_workflow_stages(self):
        rid=self.receive(qty=2)
        self.assertEqual(stock.read(self.state)['products'][self.pid]['incomingQuantity'],2)
        ids=self.register(rid,['OVERVIEW-1'])
        row=stock.read(self.state)['products'][self.pid]
        self.assertEqual((row['incomingQuantity'],row['onHandQuantity']),(1,1))
        self.act_assets('allocate',ids,client='Client',reference='PO',jobId='JOB')
        row=stock.read(self.state)['products'][self.pid]
        self.assertEqual((row['onHandQuantity'],row['allocatedQuantity']),(1,1))
        self.act_assets('dispatch',ids)
        row=stock.read(self.state)['products'][self.pid]
        self.assertEqual((row['onHandQuantity'],row['allocatedQuantity'],row['deliveredQuantity']),(0,0,1))
        rid=self.receive('Trident refurb stock',1);refurb_ids=self.state['stock']['receipts'][rid]['assetIds']
        self.assertEqual(stock.read(self.state)['products'][self.pid]['refurbingQuantity'],1)
        self.act_assets('ready',refurb_ids,technician='Rick Leef',notes='Refurb complete')
        row=stock.read(self.state)['products'][self.pid]
        self.assertEqual((row['refurbingQuantity'],row['refurbedQuantity'],row['onHandQuantity']),(0,1,0))
        self.act_assets('relabel',refurb_ids,number='OVERVIEW-REFURB',notes='New number')
        row=stock.read(self.state)['products'][self.pid]
        self.assertEqual((row['refurbedQuantity'],row['onHandQuantity']),(1,1))

    def test_dated_labels_current_priority_and_client_so_history(self):
        rid=self.receive(qty=1);ids=self.register(rid,['LABEL-A'])
        self.act_assets('relabel',ids,number='LABEL-B',notes='New refurb label',date='2026-10-15')
        self.act_assets('relabel',ids,number='LABEL-C',notes='Replacement label',date='2026-10-16')
        asset=stock.read(self.state)['assets'][ids[0]]
        self.assertEqual(asset['number'],'LABEL-C')
        self.assertEqual([{k:h[k] for k in ('number','validFrom','validTo')} for h in asset['numberHistory']],[{'number':'LABEL-A','validFrom':'2026-10-14','validTo':'2026-10-15'},{'number':'LABEL-B','validFrom':'2026-10-15','validTo':'2026-10-16'},{'number':'LABEL-C','validFrom':'2026-10-16','validTo':None}])
        before=copy.deepcopy(self.state)
        with self.assertRaises(Problem):self.act_assets('relabel',ids,number='BACKDATED',notes='Test',date='2026-10-15')
        self.assertEqual(self.state,before)
        self.act_assets('allocate',ids,client='First client',referenceType='SO',reference='SO-123',jobId='JOB-123',date='2026-10-17')
        self.act_assets('dispatch',ids,date='2026-10-18')
        self.act_assets('return',ids,condition='Ready for stock',date='2026-10-20')
        self.act_assets('allocate',ids,client='Next client',referenceType='PO',reference='PO-456',jobId='JOB-456',date='2026-10-21')
        history=stock.read(self.state)['assets'][ids[0]]['clientHistory']
        self.assertEqual(len(history),2);self.assertEqual(history[0]['referenceType'],'SO')
        self.assertEqual(history[0]['allocatedAt'],'2026-10-17');self.assertEqual(history[0]['returnedAt'],'2026-10-20')
        self.assertEqual(history[1]['client'],'Next client');self.assertNotIn('returnedAt',history[1])
        with self.assertRaises(Problem):self.act_assets('dispatch',ids,date='2026-10-19')

    def test_external_old_number_unknown_start_and_invalid_reference_type(self):
        rid=self.receive('Trident refurb stock',1);asset_id=self.state['stock']['receipts'][rid]['assetIds'][0]
        self.assertIsNone(stock.read(self.state)['assets'][asset_id]['numberHistory'][0]['validFrom'])
        rid=self.receive(qty=1);ids=self.register(rid,['TEST-SO'])
        with self.assertRaises(Problem):self.act_assets('allocate',ids,client='Client',reference='123',referenceType='Other',jobId='JOB')

    def test_bad_inputs_and_missing_technician_or_work_notes(self):
        with self.assertRaises(Problem):self.receive(qty=True)
        rid=self.receive('Trident refurb stock',1)
        for values in [{'technician':'Office','notes':'Complete'},{'technician':'Rick Leef','notes':''}]:
            with self.assertRaises(Problem):self.run_action('refurb_receipt_ready',receiptId=rid,receiptRevision=1,**values)

class StockImportTests(unittest.TestCase):
    def test_import_preserves_unknown_counts_duplicates_and_notes(self):
        candidate=stock_import.preview(workbook(),'Example.xlsx')
        self.assertEqual(len(candidate['products']),2);self.assertEqual(candidate['listedUnits'],5)
        self.assertEqual(candidate['products'][1]['sourceQuantity'],None)
        self.assertEqual(candidate['products'][0]['sourceNotes']['currentStock'],'Old asset 123 / New asset 456')
        self.assertEqual(candidate['unassignedRows'][0]['row'],4)
        original={'quotes':{'one':{'saved':True}}};state=stock_import.apply(original,candidate,'admin')
        self.assertEqual(state['quotes'],original['quotes']);self.assertEqual(state['stock']['assets'],{})
        self.assertEqual(stock.read(state)['summary']['Available stock'],5)
        self.assertEqual(sum(p['openingQuantityConfirmed'] for p in state['stock']['products'].values() if p.get('openingQuantityConfirmed')),1)
        with self.assertRaises(Problem):stock_import.apply(state,candidate,'admin')
    def test_export_retains_original_layout_unknown_quantities_and_adds_activity(self):
        raw=workbook();state=stock_import.apply({},stock_import.preview(raw),'admin',raw)
        prod=next(iter(state['stock']['products']))
        state=stock.action(state,{'action':'receive','productId':prod,'quantity':1,'pool':'New stock','client':'Supplier'},'admin')
        rid=list(state['stock']['receipts'])[-1]
        state=stock.action(state,{'action':'asset','receiptId':rid,'receiptRevision':1,'numbers':['NEW-01']},'admin')
        exported=stock_import.export(state)
        parsed=stock_import.preview(exported)
        self.assertEqual(parsed['products'][0]['sourceQuantity'],6)
        self.assertIn('Confirmed stock on hand: 1 available',parsed['products'][0]['sourceNotes']['currentStock'])
        self.assertIn('Old asset 123 / New asset 456',parsed['products'][0]['sourceNotes']['currentStock'])
        self.assertIn('current stock on hand',parsed['products'][0]['sourceNotes']['additional'])
        self.assertEqual(parsed['unassignedRows'][0]['row'],4)
        self.assertNotIn('workbookTemplate',stock.read(state))
        with zipfile.ZipFile(io.BytesIO(raw)) as old,zipfile.ZipFile(io.BytesIO(exported)) as new:
            self.assertEqual(old.read('xl/workbook.xml'),new.read('xl/workbook.xml'))
        before=copy.deepcopy(state);stock_import.export(state);self.assertEqual(state,before)

    def test_opening_labelling_and_dispatch_never_double_count(self):
        state=stock_import.apply({},stock_import.preview(workbook()),'admin',workbook())
        original=copy.deepcopy(state)
        self.assertEqual(stock_import.confirm_quantities(state,'admin'),original)
        prod=next(iter(state['stock']['products']))
        rid=state['stock']['products'][prod]['openingReceiptId']
        self.assertEqual(stock.read(state)['products'][prod]['onHandQuantity'],5)
        state=stock.action(state,{'action':'asset','receiptId':rid,'receiptRevision':1,'numbers':['OPEN-1','OPEN-2']},'admin')
        self.assertEqual(stock.read(state)['summary']['Available stock'],5)
        self.assertEqual(stock.read(state)['products'][prod]['awaitingOpeningNumbers'],3)
        key=next(iter(state['stock']['assets']))
        state=stock.action(state,{'action':'allocate','assetIds':[key],'revisions':{key:1},'client':'Client','reference':'PO','jobId':'JOB'},'admin')
        self.assertEqual(stock.read(state)['products'][prod]['onHandQuantity'],5)
        self.assertEqual(stock.read(state)['summary']['Available stock'],4)
        state=stock.action(state,{'action':'dispatch','assetIds':[key],'revisions':{key:2}},'admin')
        self.assertEqual(stock.read(state)['products'][prod]['onHandQuantity'],4)
        parsed=stock_import.preview(stock_import.export(state))
        self.assertEqual(parsed['products'][0]['sourceQuantity'],4)
        self.assertIsNone(parsed['products'][1]['sourceQuantity'])
        with self.assertRaises(Problem):stock.action(state,{'action':'asset','receiptId':rid,'receiptRevision':2,'numbers':['X1','X2','X3','X4']},'admin')

    def test_selected_stock_reports_are_read_only_and_filter_contents(self):
        from pypdf import PdfReader
        from openpyxl import load_workbook
        import base64
        state=stock_import.apply({},stock_import.preview(workbook()),'admin',workbook())
        original=copy.deepcopy(state)
        report=stock_reports.create(state,{'format':'xlsx','scope':'products'})
        sheet=load_workbook(io.BytesIO(base64.b64decode(report['file']))).active
        self.assertEqual(sheet['E2'].value,5);self.assertEqual(sheet['E3'].value,'Not recorded')
        self.assertEqual(sheet['A2'].value,'BED-MoH')
        report=stock_reports.create(state,{'format':'pdf','scope':'products'})
        pdf=PdfReader(io.BytesIO(base64.b64decode(report['file'])))
        self.assertGreater(float(pdf.pages[0].mediabox.height),float(pdf.pages[0].mediabox.width))
        self.assertIn('BED-MoH',pdf.pages[0].extract_text())
        self.assertEqual(stock_reports.tables(state,'delivered')[0][2],[])
        self.assertEqual(stock_reports.tables(state,'assets')[1][2][0][1],5)
        self.assertEqual(state,original)
        for values in [{'format':'pdf','scope':'workbook'},{'format':'exe','scope':'products'},{'format':'xlsx','scope':'anything'}]:
            with self.assertRaises(Problem):stock_reports.create(state,values)
        api=test_accounts.AccountsAPI();api.setUp()
        self.assertEqual(api.call('/api/stock-report','POST',{'format':'pdf','scope':'products'},token='tech')[0],403)
        self.assertEqual(api.call('/api/stock-report','POST',{'format':'pdf','scope':'products'})[0],200)

    def test_routes_admin_only_and_cas(self):
        api=test_accounts.AccountsAPI();api.setUp();api.app.stock_source=workbook()
        for path in ['/api/stock','/api/stock-source-preview','/api/stock-export']:
            self.assertEqual(api.call(path,token='tech')[0],403)
        self.assertEqual(api.call('/api/stock','POST',{'action':'product'},token='tech')[0],403)
        body={'useLocalSource':True,'confirmed':True}
        self.assertEqual(api.call('/api/stock-import','POST',body,match=False)[0],409)
        self.assertEqual(api.call('/api/stock-import','POST',body)[0],200)
        self.assertEqual(len(api.call('/api/stock')[1]['products']),2)
        self.assertEqual(api.call('/api/stock-import','POST',body)[0],409)
        self.assertEqual(api.call('/api/stock-export')[0],200)
        api.app.local=False
        api.app.stock_source=None
        self.assertEqual(api.call('/api/stock-export')[0],200)
        self.assertEqual(api.call('/api/stock-source-preview')[0],404)
