import base64,copy,io,json,unittest
import openpyxl
from pypdf import PdfReader
import weekly_report,weekly_export,leave

class LeaveDescriptions(unittest.TestCase):
 def setUp(self):
  self.week='2026-10-12';self.description='Bereavement leave - family loss & support <approved>'
  self.state={'staffSettings':{'technicians':[{'name':'Sample technician','rego':''}]},'quotes':{'keep':{'value':1}}}
  self.days=[{'status':'Other','other':self.description,'hours':8,'notes':'Original staff note'}]+[{'status':'Work','hours':0,'notes':''} for _ in range(6)]
  self.state['timesheets']={'old':{'technician':'Sample technician','weekStart':self.week,'submittedAt':'2026-10-12T08:00:00Z','days':[{**day,'other':'Old description'} for day in self.days]},'corrected':{'technician':'Sample technician','weekStart':self.week,'updatedAt':'2026-10-13T08:00:00Z','days':self.days}}
  self.state=weekly_report.save_notes(self.state,{'week':self.week,'notes':[{'technician':'Sample technician','date':self.week,'note':'Admin note: checked & confirmed'}]})
  self.days=self.state['timesheets']['corrected']['days']
 def book(self,result):return openpyxl.load_workbook(io.BytesIO(base64.b64decode(result['file'])))
 def text(self,result):return '\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(base64.b64decode(result['pdf']))).pages)
 def approved(self):
  state,request=leave.create(self.state,{'technician':'Sample technician','leaveType':'Other','other':self.description,'startDate':self.week,'endDate':self.week,'hoursPerDay':8,'reason':'PRIVATE REQUEST REASON'})
  return leave.decide(state,{'id':request['id'],'revision':request['revision'],'status':'Approved'},'admin@example.nz')[0]
 def test_latest_description_and_admin_note_are_independent(self):
  before=copy.deepcopy(self.state);day=weekly_report.week_rows(self.state,self.week)[0]['days'][0]
  self.assertEqual(day['otherDescription'],self.description);self.assertEqual(day['note'],'Admin note: checked & confirmed');self.assertEqual(day['all'],8)
  self.assertEqual(self.state,before)
 def test_approved_description_without_submission_and_private_reason_excluded(self):
  state=self.approved();state.pop('timesheets')
  rows=weekly_report.week_rows(state,self.week);self.assertEqual(rows[0]['days'][0]['otherDescription'],self.description)
  self.assertNotIn('PRIVATE REQUEST REASON',json.dumps(rows));self.assertEqual(rows[0]['days'][0]['other'],8)
  self.assertIn(self.description,self.text(weekly_report.pdf(state,{'week':self.week})))
 def test_current_approved_description_overrides_stale_submission(self):
  state=self.approved();state['timesheets']['corrected']['days'][0]['other']='Stale description'
  self.assertEqual(weekly_report.week_rows(state,self.week)[0]['days'][0]['otherDescription'],self.description)
 def test_weekly_excel_visible_leave_and_admin_notes_cells_both_modes(self):
  before=copy.deepcopy(self.state)
  for mode,payable in [('all',8),('work',0)]:
   book=self.book(weekly_export.weekly(self.state,{'week':self.week,'mode':mode}));summary=book['Weekly hours'];detail=book['Daily detail']
   self.assertEqual(summary['C2'].value,payable);self.assertEqual(summary['C2'].data_type,'n')
   self.assertEqual(summary['B4'].value,'Leave / notes')
   for text in ('Other 8h',self.description,'Admin note: checked & confirmed'):self.assertIn(text,summary['C4'].value)
   self.assertEqual(detail['J2'].value,'Admin note: checked & confirmed');self.assertEqual(detail['K2'].value,self.description)
   self.assertEqual(detail['F2'].value,8);self.assertEqual(detail['K1'].value,'Other leave description')
   self.assertEqual(summary['C4'].data_type,'s');self.assertTrue(summary['C4'].alignment.wrap_text)
   total=next(r for r in summary.iter_rows(values_only=True) if r[0]=='Total' and r[1]=='Payable');self.assertEqual(total[2],payable)
  self.assertEqual(self.state,before)
 def test_sick_annual_and_zero_hour_notes_are_visible_in_main_excel_cells(self):
  self.days[1]={'status':'Annual leave','hours':8,'notes':'Annual note'};self.days[2]={'status':'Sick leave','hours':4,'notes':''};self.days[3]={'status':'Work','hours':0,'notes':'Zero hours note'}
  sheet=self.book(weekly_export.weekly(self.state,{'week':self.week}))['Weekly hours']
  self.assertIn('Annual 8h',sheet['D4'].value);self.assertIn('Annual note',sheet['D4'].value);self.assertEqual(sheet['E4'].value,'Sick 4h');self.assertEqual(sheet['F4'].value,'Zero hours note')
 def test_weekly_and_selected_period_pdfs_keep_description_and_totals(self):
  for mode in ('all','work'):
   weekly=self.text(weekly_report.pdf(self.state,{'week':self.week,'mode':mode}));period=self.text(weekly_report.period_pdf(self.state,{'start':self.week,'end':self.week,'mode':mode}))
   for text in (weekly,period):
    self.assertIn('Other leave descriptions',text);self.assertIn(self.description,text);self.assertIn('12/10/2026',text);self.assertNotIn('Old description',text)
   self.assertIn('Admin note: checked & confirmed',weekly)
 def test_selected_period_excel_has_descriptions_for_only_selected_weeks(self):
  book=self.book(weekly_export.period(self.state,{'start':self.week,'end':self.week}));rows=list(book['Other leave details'].values)
  self.assertEqual(rows[1],('Sample technician',self.week,8,self.description))
  annotations=list(book['Daily notes and leave'].values);self.assertIn('Admin note: checked & confirmed',annotations[1][2]);self.assertIn('Other 8h',annotations[1][2])
  other=self.book(weekly_export.period(self.state,{'start':'2026-10-19','end':'2026-10-19'}));self.assertEqual(other['Other leave details'].max_row,1)
 def test_descriptions_and_admin_notes_are_plain_text_in_excel(self):
  self.days[0]['other']='=HYPERLINK("https://example.nz")'
  self.state['weeklyReportNotes'][self.week]['sample technician'][self.week]='=1+1'
  book=self.book(weekly_export.weekly(self.state,{'week':self.week}));detail=book['Daily detail']
  for address in ('J2','K2'):self.assertEqual(detail[address].data_type,'s')
  self.assertEqual(detail['J2'].value,'=1+1');self.assertEqual(detail['K2'].value,self.days[0]['other'])

if __name__=='__main__':unittest.main()
