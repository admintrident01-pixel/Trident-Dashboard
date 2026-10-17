import base64,copy,io,unittest
from pathlib import Path
from openpyxl import load_workbook
from pypdf import PdfReader
import stock,stock_take,stock_reports,test_stock_standard

class StockTakeTests(unittest.TestCase):
    def setUp(self):
        fixture=test_stock_standard.StandardReportTests();fixture.setUp();self.state=fixture.state;self.pid=fixture.pid
        available=[a for a in self.state['stock']['assets'].values() if a['status']=='Available']
        available[0].update(status='Refurbishing',location='Workshop')
        available[1].update(status='Retired')
    def test_counts_match_physical_checklist_and_do_not_mutate_stock(self):
        before=copy.deepcopy(self.state);quantities,assets=stock_take.rows(self.state)
        self.assertEqual(quantities[0][4],8);self.assertEqual(len(quantities[0]),8);self.assertEqual(len(quantities[0][3].splitlines()),8);self.assertEqual(quantities[0][5],'');self.assertEqual(quantities[0][6:],['',''])
        self.assertEqual(len(assets),8)
        self.assertTrue(all(a[5]=='[ ]' and a[6]=='' for a in assets))
        self.assertFalse(any('With client' in a[4] or 'Retired' in a[4] for a in assets))
        self.assertEqual(self.state,before)
    def test_unknown_quantity_is_not_assumed_zero(self):
        self.state=stock.action(self.state,{'action':'product','code':'UNKNOWN','description':'Unconfirmed product'},'test')
        p=next(p for p in self.state['stock']['products'].values() if p['code']=='UNKNOWN');p.update(quantityBasis='listed_assets',sourceQuantity=None)
        q,a=stock_take.rows(self.state);self.assertEqual(next(r for r in q if r[0]=='UNKNOWN')[4],'Not recorded')
        self.assertFalse(any(r[0]=='UNKNOWN' for r in a))
    def test_excel_has_blank_count_fields_repeat_headers_and_print_setup(self):
        b=load_workbook(io.BytesIO(stock_take.excel(self.state,'2026-10-15')))
        self.assertEqual(b.sheetnames,['Quantity count','Asset checklist'])
        for s in b:
            self.assertEqual(s.page_setup.orientation,'landscape');self.assertEqual(s.page_setup.fitToWidth,1)
            self.assertEqual(s.print_title_rows,'$1:$7');self.assertTrue(s.print_area)
        self.assertIn(b['Quantity count']['F8'].value,(None,''));self.assertTrue(b['Quantity count']['D8'].value);self.assertEqual(b['Quantity count']['D8'].alignment.textRotation,45);self.assertIn(b['Quantity count']['G8'].value,(None,''));self.assertIn(b['Quantity count']['H8'].value,(None,''));self.assertEqual([c.value for c in b['Quantity count'][7]],stock_take.QHEAD)
        self.assertIn('Technician:',b['Quantity count']['A3'].value)
        self.assertEqual(b['Asset checklist']['D8'].value,'[ ]')
    def test_pdf_and_api_have_manual_check_sections_and_blank_additions(self):
        result=stock_reports.create(self.state,{'scope':'stock_take','format':'pdf'})
        self.assertTrue(result['name'].startswith('Stock take Sheet-'))
        pdf=PdfReader(io.BytesIO(base64.b64decode(result['file'])));text='\n'.join(p.extract_text() for p in pdf.pages)
        self.assertTrue(all(p.mediabox.width>p.mediabox.height for p in pdf.pages))
        for label in ('Quantity count','Asset checklist','Unlisted assets / corrections','Technician:','Counted','Expected','Asset numbers','Difference','Notes'):self.assertIn(label,' '.join(text.split()))
        self.assertIn('stock_take_template.xlsx',(Path(stock_take.__file__).parent/'Dockerfile').read_text())
