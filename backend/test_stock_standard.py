import base64,copy,io,unittest
from pathlib import Path
from openpyxl import load_workbook
from pypdf import PdfReader
import stock,stock_standard,stock_reports

class StandardReportTests(unittest.TestCase):
    def setUp(self):
        self.state=stock.action({},dict(action='product',code='TEST-NEW',condition='NEW',description='Test mattress'),'test')
        self.pid=next(iter(self.state['stock']['products']))
        self.state=stock.action(self.state,dict(action='receive_stock',date='2026-10-15',productId=self.pid,pool='New stock',quantity=10,numbers=['ASSET-'+str(i) for i in range(10)],reference='SUP-PO'),'test')
        ids=list(self.state['stock']['assets'])
        self.move('allocate',ids[:2],client='Test client',referenceType='SO',reference='SO-1',jobId='JOB-1')
        self.move('dispatch',ids[1:2])
        self.move('relabel',ids[2:3],number='REPLACEMENT',notes='Replace label')
        self.state=stock.action(self.state,dict(action='stock_settings',minimums={self.pid:12}),'test')
    def move(self,action,ids,**values):
        self.state=stock.action(self.state,dict(action=action,date='2026-10-15',assetIds=ids,revisions={key:self.state['stock']['assets'][key]['revision'] for key in ids},**values),'test')
    def test_live_quantities_chart_and_asset_grid_reconcile_without_mutation(self):
        before=copy.deepcopy(self.state);raw=stock_standard.excel(self.state,'2026-10-15');book=load_workbook(io.BytesIO(raw))
        self.assertEqual(book.sheetnames,['Stock overview','Stock quantities','Asset numbers'])
        row=[book['Stock quantities'].cell(7,c).value for c in range(1,12)]
        self.assertEqual(row[5:],[8,1,9,1,12,4])
        chart=book['Stock overview']._charts[0]
        self.assertEqual(chart.type,'bar');self.assertEqual(chart.grouping,'stacked')
        self.assertEqual([sum(pt.v for pt in ser.val.numRef.numCache.pt) for ser in chart.series],[9,0,0])
        assets=book['Asset numbers'];self.assertEqual(assets['D7'].value,10)
        self.assertIsNotNone(assets['N7'].value)
        text='\n'.join(str(c.value or '') for row in assets for c in row)
        self.assertIn('REPLACEMENT',text);self.assertIn('Earlier: ASSET-2',text);self.assertIn('Status: With client',text)
        self.assertEqual(self.state,before)
    def test_pdf_is_landscape_with_chart_and_quantities_only(self):
        raw=stock_standard.pdf(self.state,'2026-10-15');pdf=PdfReader(io.BytesIO(raw));texts=[p.extract_text() for p in pdf.pages]
        self.assertTrue(all(float(p.mediabox.width)>float(p.mediabox.height) for p in pdf.pages))
        self.assertEqual(len(pdf.pages),2)
        self.assertIn('Stock on hand by product and condition',texts[0]);self.assertIn('Stock quantities',texts[1])
        text='\n'.join(texts)
        self.assertNotIn('Asset numbers by stock',text)
        self.assertNotIn('REPLACEMENT',text)
        self.assertNotIn('ASSET-2',text)
    def test_chart_orders_by_product_code_instead_of_quantity(self):
        for code,count in [('100200',5),('100100',1)]:
            self.state=stock.action(self.state,dict(action='product',code=code,description='Sorting test'),'test')
            key=next(p['id'] for p in self.state['stock']['products'].values() if p['code']==code)
            self.state=stock.action(self.state,dict(action='receive_stock',date='2026-10-15',productId=key,pool='New stock',quantity=count,numbers=[code+'-'+str(i) for i in range(count)]),'test')
        self.assertEqual([g['code'] for g in stock_standard.snapshot(self.state)['chart']],['100100','100200','TEST-NEW'])

    def test_unconfirmed_quantity_is_not_exported_as_zero(self):
        self.state=stock.action(self.state,dict(action='product',code='UNKNOWN',description='Unknown quantity'),'test')
        key=next(k for k in self.state['stock']['products'] if k!=self.pid)
        self.state['stock']['products'][key].update(quantityBasis='listed_assets',sourceQuantity=None)
        snap=stock_standard.snapshot(self.state);row=next(r for r in snap['quantities'] if r[0]=='UNKNOWN')
        self.assertEqual(row[5],'Not recorded');self.assertEqual(row[7],'Not recorded')
        self.assertEqual(len(snap['chart']),1)
    def test_api_report_formats_and_container_template_are_available(self):
        for fmt in ('xlsx','pdf'):
            report=stock_reports.create(self.state,dict(format=fmt,scope='standard_soh'))
            self.assertTrue(report['name'].startswith(stock_standard.TITLE));self.assertTrue(base64.b64decode(report['file']))
        self.assertIn('standard_soh_template.xlsx',(Path(stock_standard.__file__).parent/'Dockerfile').read_text())
        self.assertTrue(Path(stock_standard.__file__).with_name('standard_soh_template.xlsx').exists())
